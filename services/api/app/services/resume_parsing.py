from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass
from pathlib import PurePath

from docx import Document
from pypdf import PdfReader


class ResumeFileError(ValueError):
    """Raised when an uploaded resume is invalid or unsupported."""


@dataclass(frozen=True)
class ParsedResumeFile:
    filename: str
    extension: str
    raw_text: str


class ResumeFileParser:
    allowed_extensions = {".pdf", ".docx", ".txt", ".md", ".markdown"}
    allowed_mime_types = {
        ".pdf": {"application/pdf", "application/octet-stream"},
        ".docx": {
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "application/octet-stream",
        },
        ".txt": {"text/plain", "application/octet-stream"},
        ".md": {"text/markdown", "text/plain", "application/octet-stream"},
        ".markdown": {"text/markdown", "text/plain", "application/octet-stream"},
    }

    def parse(
        self, *, filename: str | None, content_type: str | None, data: bytes
    ) -> ParsedResumeFile:
        safe_filename = self._safe_filename(filename)
        extension = PurePath(safe_filename).suffix.lower()
        if extension not in self.allowed_extensions:
            raise ResumeFileError("Only PDF, DOCX, Markdown, and TXT resumes are supported")

        normalized_type = (content_type or "application/octet-stream").split(";", 1)[0].lower()
        if normalized_type not in self.allowed_mime_types[extension]:
            raise ResumeFileError(f"Unsupported MIME type for {extension}: {normalized_type}")

        if extension == ".pdf":
            raw_text = self._parse_pdf(data)
        elif extension == ".docx":
            raw_text = self._parse_docx(data)
        else:
            raw_text = self._parse_text(data)

        normalized_text = self._normalize_text(raw_text)
        if not normalized_text:
            raise ResumeFileError("The resume did not contain extractable text")
        return ParsedResumeFile(safe_filename, extension, normalized_text)

    @staticmethod
    def _safe_filename(filename: str | None) -> str:
        if not filename:
            raise ResumeFileError("A filename is required")
        normalized = filename.replace("\\", "/").rsplit("/", 1)[-1].strip()
        if normalized in {"", ".", ".."}:
            raise ResumeFileError("Invalid filename")
        return normalized

    @staticmethod
    def _parse_pdf(data: bytes) -> str:
        if not data.lstrip().startswith(b"%PDF"):
            raise ResumeFileError("The uploaded file is not a valid PDF")
        try:
            reader = PdfReader(io.BytesIO(data))
            return "\n".join(page.extract_text() or "" for page in reader.pages)
        except Exception as exc:
            raise ResumeFileError("Unable to parse the PDF resume") from exc

    @staticmethod
    def _parse_docx(data: bytes) -> str:
        if not zipfile.is_zipfile(io.BytesIO(data)):
            raise ResumeFileError("The uploaded file is not a valid DOCX")
        try:
            document = Document(io.BytesIO(data))
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            table_cells = [
                cell.text for table in document.tables for row in table.rows for cell in row.cells
            ]
            return "\n".join(paragraphs + table_cells)
        except Exception as exc:
            raise ResumeFileError("Unable to parse the DOCX resume") from exc

    @staticmethod
    def _parse_text(data: bytes) -> str:
        if b"\x00" in data:
            raise ResumeFileError("The text resume contains binary data")
        try:
            return data.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ResumeFileError("Text resumes must use UTF-8 encoding") from exc

    @staticmethod
    def _normalize_text(text: str) -> str:
        lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.splitlines()]
        return "\n".join(line for line in lines if line)
