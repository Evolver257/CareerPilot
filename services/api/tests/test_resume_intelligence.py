from io import BytesIO

from docx import Document
from httpx import AsyncClient
from reportlab.pdfgen import canvas

from app.services.resume_extraction import HeuristicResumeExtractor
from app.services.resume_parsing import ResumeFileParser


def _docx_bytes() -> bytes:
    document = Document()
    document.add_heading("Resume", level=1)
    document.add_paragraph("Skills")
    document.add_paragraph("Python, FastAPI, RAG")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _pdf_bytes() -> bytes:
    output = BytesIO()
    pdf = canvas.Canvas(output)
    pdf.drawString(72, 720, "AI Agent Engineer - Python RAG")
    pdf.save()
    return output.getvalue()


def test_pdf_and_docx_parsers_extract_text() -> None:
    parser = ResumeFileParser()

    pdf = parser.parse(filename="resume.pdf", content_type="application/pdf", data=_pdf_bytes())
    docx = parser.parse(
        filename="resume.docx",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        data=_docx_bytes(),
    )

    assert "AI Agent Engineer" in pdf.raw_text
    assert "Python, FastAPI, RAG" in docx.raw_text


def test_narrative_skill_section_produces_canonical_skill_tags() -> None:
    profile = HeuristicResumeExtractor().extract(
        """张三

专业技能
1. 熟悉 Java 编程语言和 SpringBoot 后端框架，掌握 MySQL 与 PostgreSQL。
2. 熟悉 MyBatis-Plus、Redis、RabbitMQ、Linux 与 Git。
成都信息工程大学
"""
    )

    assert {"Java", "Spring Boot", "MySQL", "PostgreSQL", "MyBatis", "Redis"}.issubset(
        profile.skills
    )
    assert "成都信息工程大学" not in profile.skills
    assert all("熟悉" not in skill for skill in profile.skills)


def test_pdf_lines_are_grouped_into_education_and_project_entries() -> None:
    profile = HeuristicResumeExtractor().extract(
        """邱铭曦
教育经历
2026.09-2029.06 华中科技大学 软件工程（硕士研究生）
2022.09-2026.06 软件工程（本科）
GPA：3.59 / 5.0 专业排名：17/177
专业技能
1. 熟悉 Java、SpringBoot、MySQL 与 Redis。
项目经历
基于可拓数据挖掘的学生成绩数据分析平台 2024/09-2024/12
负责任务: 外卖平台订单模块、购物车模块的后台开发
1. 使用 SpringBoot、Redis 与 MySQL 完成订单模块开发。
竞赛奖项
负责任务: 数据分析、模型训练、前端可视化
1. 使用 Pandas 完成数据清洗。
2. 使用 SVM 完成模型训练，准确度达到 97.2%。
1. 中国高校计算机大赛全国二等奖
2. 蓝桥杯 Java 软件开发四川赛区二等奖
2024/04-2024/07 校园生活外卖平台开发 后端开发
"""
    )

    assert len(profile.education) == 2
    assert profile.education[0].institution == "华中科技大学"
    assert profile.education[0].degree == "硕士研究生"
    assert profile.education[0].start_date == "2026.09"
    assert profile.education[1].field == "软件工程"
    assert profile.education[1].details == ["GPA：3.59 / 5.0 专业排名：17/177"]

    assert len(profile.projects) == 2
    assert profile.projects[0].name == "基于可拓数据挖掘的学生成绩数据分析平台"
    assert "数据清洗" in profile.projects[0].description
    assert "订单模块" not in profile.projects[0].description
    assert profile.projects[1].name == "校园生活外卖平台开发 后端开发"
    assert "订单模块" in profile.projects[1].description
    assert len(profile.awards) == 2


def test_compound_headings_and_pipe_titles_create_project_groups() -> None:
    profile = HeuristicResumeExtractor().extract(
        """教育背景
2022.09-2026.06 成都信息工程大学 软件工程（本科）
GPA 3.52 / 5.0 | 专业排名 14/177（前 8%）
研究与项目经历
MiniMind | 大语言模型训练与 Agent 开发
基于 PyTorch 实现 Decoder-only Transformer。
搭建 Agentic RL 训练流程。
Coding Agent | Agent Harness 设计
设计基于 Anthropic Messages API 的 Agent Loop。
眼底影像智能识别系统 | 深度学习与计算机视觉
基于改进 ResNet50 模型完成图像识别任务。
竞赛与荣誉
中国高校计算机大赛网络技术挑战赛 全国二等奖
专业技能
Python, PyTorch, FastAPI, Agent
"""
    )

    assert len(profile.education) == 1
    assert len(profile.projects) == 3
    assert [project.name for project in profile.projects] == [
        "MiniMind",
        "Coding Agent",
        "眼底影像智能识别系统",
    ]
    assert "PyTorch" in profile.projects[0].technologies
    assert len(profile.awards) == 1


async def test_resume_upload_parse_and_chunk_endpoints(client: AsyncClient) -> None:
    resume_text = """Alex Chen

Summary
AI Agent Engineer with Python, RAG, and FastAPI experience.

Projects
CareerPilot - built a resume intelligence service with PostgreSQL and Docker.

Skills
Python, FastAPI, RAG, PostgreSQL, Docker
"""
    response = await client.post(
        "/api/resumes",
        files={"file": ("alex_resume.txt", resume_text.encode("utf-8"), "text/plain")},
        data={"name": "Alex Resume"},
    )

    assert response.status_code == 201
    payload = response.json()
    assert payload["name"] == "Alex Resume"
    assert payload["raw_text"].startswith("Alex Chen")
    assert "Python" in payload["structured_profile"]["skills"]
    assert payload["chunk_count"] >= 2

    resume_id = payload["id"]
    chunks = await client.get(f"/api/resumes/{resume_id}/chunks")
    assert chunks.status_code == 200
    assert chunks.json()[0]["embedding_dimensions"] == 1024

    parsed = await client.post(f"/api/resumes/{resume_id}/parse")
    assert parsed.status_code == 200
    assert parsed.json()["chunk_count"] == payload["chunk_count"]

    listed = await client.get("/api/resumes")
    assert listed.status_code == 200
    assert listed.json()["total"] == 1
